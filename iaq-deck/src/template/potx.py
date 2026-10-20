"""Turn the emitted template pages into a PowerPoint template.

Each empty layout becomes a real slide layout in the master, with its text
boxes converted to placeholders: the client picks the layout from New Slide,
clicks a box and types, and the type, colour and position are already IAQ's.
"""
import copy, io, json, os, re, shutil, zipfile
from lxml import etree
from pptx import Presentation
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
PAGES = os.path.join(HERE, 'IAQ-Template-Pages.pptx')
INDEX = '/home/user/Brand-Method-Website/iaq-deck/template/layouts.json'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
NS = {'a': A, 'p': P, 'r': R}
q = lambda t: '{%s}%s' % (NS[t.split(':')[0]], t.split(':')[1])

# IAQ's palette, so the colour picker in PowerPoint offers the right reds
SCHEME = [('dk1', '101A2B'), ('lt1', 'FFFFFF'), ('dk2', '1E2A3D'), ('lt2', 'F6F7F9'),
          ('accent1', 'EC2326'), ('accent2', '101A2B'), ('accent3', '1E2A3D'),
          ('accent4', '64707A'), ('accent5', 'D5DBE2'), ('accent6', 'E6EAEF'),
          ('hlink', 'EC2326'), ('folHlink', '64707A')]


def layout_xml(sp_tree_children, name):
    """Wrap a converted shape tree in a slide layout part."""
    root = etree.Element(q('p:sldLayout'), nsmap={'a': A, 'r': R, 'p': P})
    root.set('type', 'obj')
    root.set('preserve', '1')
    csld = etree.SubElement(root, q('p:cSld'))
    csld.set('name', name)
    tree = etree.SubElement(csld, q('p:spTree'))
    nv = etree.SubElement(tree, q('p:nvGrpSpPr'))
    c = etree.SubElement(nv, q('p:cNvPr')); c.set('id', '1'); c.set('name', '')
    etree.SubElement(nv, q('p:cNvGrpSpPr'))
    etree.SubElement(nv, q('p:nvPr'))
    grp = etree.SubElement(tree, q('p:grpSpPr'))
    xfrm = etree.SubElement(grp, q('a:xfrm'))
    for tag in ('a:off', 'a:chOff'):
        e = etree.SubElement(xfrm, q(tag)); e.set('x', '0'); e.set('y', '0')
    for tag in ('a:ext', 'a:chExt'):
        e = etree.SubElement(xfrm, q(tag)); e.set('cx', '0'); e.set('cy', '0')
    for sp in sp_tree_children:
        tree.append(sp)
    ovr = etree.SubElement(root, q('p:clrMapOvr'))
    etree.SubElement(ovr, q('a:masterClrMapping'))
    return root


def slide_number(sp, sid):
    """The folio is a field, not a typed figure, so pages number themselves."""
    to_placeholder(sp, 12, 'sldNum', sid, 'Slide Number')
    body = sp.find(q('p:txBody'))
    for para in body.findall(q('a:p')):
        for r in para.findall(q('a:r')):
            para.remove(r)
        fld = etree.SubElement(para, q('a:fld'))
        fld.set('id', '{B7D3B2A1-4E6C-41E2-9B53-2A4C6E1F7A10}')
        fld.set('type', 'slidenum')
        rpr = etree.SubElement(fld, q('a:rPr')); rpr.set('lang', 'en-GB')
        t = etree.SubElement(fld, q('a:t')); t.text = '2'
        break
    return sp


def to_placeholder(sp, idx, kind, sid, label):
    """A text box becomes a placeholder that keeps the design's type."""
    nv = sp.find(q('p:nvSpPr'))
    cnv = nv.find(q('p:cNvPr'))
    cnv.set('id', str(sid)); cnv.set('name', label)
    csp = nv.find(q('p:cNvSpPr'))
    if 'txBox' in csp.attrib:
        del csp.attrib['txBox']
    if csp.find(q('a:spLocks')) is None:
        lk = etree.SubElement(csp, q('a:spLocks')); lk.set('noGrp', '1')
    nvpr = nv.find(q('p:nvPr'))
    ph = etree.SubElement(nvpr, q('p:ph'))
    ph.set('type', kind)
    if kind != 'title':
        ph.set('idx', str(idx))

    # carry the first run's type into the placeholder's own list style, so what
    # the client types comes out looking like what the layout shows
    body = sp.find(q('p:txBody'))
    para = body.find(q('a:p'))
    run = para.find(q('a:r')) if para is not None else None
    lst = body.find(q('a:lstStyle'))
    if lst is None:
        lst = etree.Element(q('a:lstStyle'))
        body.insert(1, lst)
    lvl = etree.SubElement(lst, q('a:lvl1pPr'))
    ppr = para.find(q('a:pPr')) if para is not None else None
    if ppr is not None:
        if ppr.get('algn'):
            lvl.set('algn', ppr.get('algn'))
        ln = ppr.find(q('a:lnSpc'))
        if ln is not None:
            lvl.append(copy.deepcopy(ln))
    bu = etree.SubElement(lvl, q('a:buNone'))
    if run is not None:
        rpr = run.find(q('a:rPr'))
        if rpr is not None:
            d = copy.deepcopy(rpr)
            d.tag = q('a:defRPr')
            for k in ('dirty', 'err', 'lang', 'smtClean'):
                d.attrib.pop(k, None)
            lvl.append(d)
    return sp


FOLIO = []


def build():
    names = [x['name'] for x in json.load(io.open(INDEX, encoding='utf-8'))]
    src = zipfile.ZipFile(PAGES)
    pres = etree.fromstring(src.read('ppt/presentation.xml'))
    prels = etree.fromstring(src.read('ppt/_rels/presentation.xml.rels'))
    order = []
    for sid in pres.find(q('p:sldIdLst')):
        rid = sid.get(q('r:id'))
        for rel in prels:
            if rel.get('Id') == rid:
                order.append('ppt/' + rel.get('Target'))
    empties = [order[i] for i in range(2, len(order), 2)]      # the empty twin of each pair
    assert len(empties) == len(names), (len(empties), len(names))

    layouts, media = [], {}
    for n, (path, name) in enumerate(zip(empties, names), 1):
        slide = etree.fromstring(src.read(path))
        rels = etree.fromstring(src.read(path.replace('slides/', 'slides/_rels/') + '.rels'))
        rmap = {rel.get('Id'): rel.get('Target') for rel in rels}
        tree = slide.find(q('p:cSld')).find(q('p:spTree'))
        kids = [k for k in tree if k.tag in (q('p:sp'), q('p:pic'))]

        # the biggest line on an untitled layout is its title; elsewhere it is
        # the page title the header carries
        texts = []
        for k in kids:
            if k.tag != q('p:sp'):
                continue
            t = ''.join(e.text or '' for e in k.iter(q('a:t')))
            sz = max([int(e.get('sz', 0)) for e in k.iter() if e.get('sz')] or [0])
            if t.strip():
                texts.append((k, t.strip(), sz))
        title_sp = None
        for k, t, sz in texts:
            if t == 'Page title':
                title_sp = k
                break
        if title_sp is None and texts:
            title_sp = max(texts, key=lambda x: x[2])[0]

        out, idx, sid = [], 20, 2
        newrels = ['<Relationship Id="rId1" Type="%s/slideMaster" Target="../slideMasters/slideMaster1.xml"/>' % R]
        rid = 2
        for k in kids:
            sid += 1
            if k.tag == q('p:pic'):
                blip = k.find('.//' + q('a:blip'))
                if blip is not None:
                    tgt = rmap[blip.get(q('r:embed'))]
                    name_in_zip = os.path.normpath(os.path.join('ppt/slideLayouts', tgt))
                    media[name_in_zip] = tgt
                    blip.set(q('r:embed'), 'rId%d' % rid)
                    newrels.append('<Relationship Id="rId%d" Type="%s/image" Target="%s"/>'
                                   % (rid, R, tgt))
                    rid += 1
                k.find(q('p:nvPicPr')).find(q('p:cNvPr')).set('id', str(sid))
                out.append(k)
                continue
            t = ''.join(e.text or '' for e in k.iter(q('a:t'))).strip()
            if not t:
                k.find(q('p:nvSpPr')).find(q('p:cNvPr')).set('id', str(sid))
                out.append(k)
                continue
            box = k.find(q('p:spPr')).find(q('a:xfrm'))
            off = box.find(q('a:off')) if box is not None else None
            x = int(off.get('x')) if off is not None else 0
            y = int(off.get('y')) if off is not None else 0
            if t.isdigit() and y > 6000000 and x < 2000000:
                FOLIO.append(slide_number(copy.deepcopy(k), 9))
                continue
            elif k is title_sp:
                out.append(to_placeholder(k, 0, 'title', sid, 'Title'))
            else:
                out.append(to_placeholder(k, idx, 'body', sid, 'Text %d' % (idx - 19)))
                idx += 1
        layouts.append((name, layout_xml(out, name), newrels))
    return src, layouts, media


if __name__ == '__main__':
    src, layouts, media = build()
    print('%d layouts' % len(layouts))
    for name, xml, rels in layouts[:3]:
        ph = len(xml.findall('.//' + q('p:ph')))
        print('  %-34s %2d shapes, %d placeholders, %d rels'
              % (name, len(xml.find(q('p:cSld')).find(q('p:spTree'))) - 2, ph, len(rels)))


# ---------------------------------------------------------------- assembly
CT_LAYOUT = 'application/vnd.openxmlformats-officedocument.presentationml.slideLayout+xml'
CT_TEMPLATE = 'application/vnd.openxmlformats-officedocument.presentationml.template.main+xml'
CT_DECK = 'application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml'


def base_package(tmp):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    prs.save(tmp)
    return tmp


def theme_iaq(xml):
    """Put IAQ's colours and faces in the theme, so the pickers offer them."""
    t = etree.fromstring(xml)
    scheme = t.find('.//' + q('a:clrScheme'))
    for (slot, hexv), node in zip(SCHEME, scheme):
        node.clear()
        node.tag = q('a:' + slot)
        if slot in ('dk1', 'lt1'):
            e = etree.SubElement(node, q('a:sysClr'))
            e.set('val', 'windowText' if slot == 'dk1' else 'window')
            e.set('lastClr', hexv)
        else:
            e = etree.SubElement(node, q('a:srgbClr')); e.set('val', hexv)
    fonts = t.find('.//' + q('a:fontScheme'))
    for scope, face in (('a:majorFont', 'Poppins SemiBold'), ('a:minorFont', 'Urbanist')):
        blk = fonts.find(q(scope))
        for tag in ('a:latin', 'a:ea', 'a:cs'):
            e = blk.find(q(tag))
            if e is None:
                e = etree.SubElement(blk, q(tag))
            e.set('typeface', face if tag == 'a:latin' else '')
    fonts.set('name', 'IAQ')
    scheme.set('name', 'IAQ')
    return etree.tostring(t, xml_declaration=True, encoding='UTF-8', standalone=True)


def master_xml(xml, n):
    m = etree.fromstring(xml)
    tree = m.find(q('p:cSld')).find(q('p:spTree'))
    for sp in list(tree.findall(q('p:sp'))):
        ph = sp.find('.//' + q('p:ph'))
        if ph is not None and ph.get('type') in ('dt', 'ftr', 'sldNum'):
            tree.remove(sp)
    if FOLIO:
        tree.append(copy.deepcopy(FOLIO[0]))
    lst = m.find(q('p:sldLayoutIdLst'))
    for child in list(lst):
        lst.remove(child)
    for i in range(n):
        e = etree.SubElement(lst, q('p:sldLayoutId'))
        e.set('id', str(2147483649 + i))
        e.set(q('r:id'), 'rId%d' % (i + 1))
    # the paper the deck is set on, in case a layout ever shows through
    cs = m.find(q('p:cSld'))
    if cs.find(q('p:bg')) is None:
        bg = etree.Element(q('p:bg'))
        pr = etree.SubElement(bg, q('p:bgPr'))
        fill = etree.SubElement(pr, q('a:solidFill'))
        clr = etree.SubElement(fill, q('a:srgbClr')); clr.set('val', 'F6F7F9')
        etree.SubElement(pr, q('a:effectLst'))
        cs.insert(0, bg)
    return etree.tostring(m, xml_declaration=True, encoding='UTF-8', standalone=True)


def write_potx(out):
    src, layouts, media = build()
    tmp = os.path.join(HERE, '_base.pptx')
    base_package(tmp)
    zb = zipfile.ZipFile(tmp)
    skip = re.compile(r'^ppt/slideLayouts/')
    master_rels = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">']
    for i, (name, _, _) in enumerate(layouts, 1):
        master_rels.append('<Relationship Id="rId%d" Type="%s/slideLayout" '
                           'Target="../slideLayouts/slideLayout%d.xml"/>' % (i, R, i))
    for rel in etree.fromstring(zb.read('ppt/slideMasters/_rels/slideMaster1.xml.rels')):
        if rel.get('Type').endswith('/theme'):
            master_rels.append('<Relationship Id="rId%d" Type="%s" Target="%s"/>'
                               % (len(layouts) + 1, rel.get('Type'), rel.get('Target')))
    master_rels.append('</Relationships>')

    ct = zb.read('[Content_Types].xml').decode('utf-8')
    ct = re.sub(r'<Override PartName="/ppt/slideLayouts/[^"]+"[^>]*/>', '', ct)
    ct = ct.replace(CT_DECK, CT_TEMPLATE)
    adds = ''.join('<Override PartName="/ppt/slideLayouts/slideLayout%d.xml" ContentType="%s"/>'
                   % (i, CT_LAYOUT) for i in range(1, len(layouts) + 1))
    for ext, mime in (('png', 'image/png'), ('jpeg', 'image/jpeg'), ('jpg', 'image/jpeg'),
                      ('fntdata', 'application/x-fontdata')):
        if 'Extension="%s"' % ext not in ct:
            adds += '<Default Extension="%s" ContentType="%s"/>' % (ext, mime)
    ct = ct.replace('</Types>', adds + '</Types>')

    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        for n in zb.namelist():
            if skip.match(n):
                continue
            if n == '[Content_Types].xml':
                z.writestr(n, ct)
            elif n == 'ppt/slideMasters/slideMaster1.xml':
                z.writestr(n, master_xml(zb.read(n), len(layouts)))
            elif n == 'ppt/slideMasters/_rels/slideMaster1.xml.rels':
                z.writestr(n, ''.join(master_rels))
            elif n.startswith('ppt/theme/theme1.xml'):
                z.writestr(n, theme_iaq(zb.read(n)))
            else:
                z.writestr(n, zb.read(n))
        for i, (name, xml, rels) in enumerate(layouts, 1):
            z.writestr('ppt/slideLayouts/slideLayout%d.xml' % i,
                       etree.tostring(xml, xml_declaration=True, encoding='UTF-8',
                                      standalone=True))
            z.writestr('ppt/slideLayouts/_rels/slideLayout%d.xml.rels' % i,
                       '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                       '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
                       'relationships">' + ''.join(rels) + '</Relationships>')
        for zip_name, target in media.items():
            inner = os.path.normpath(os.path.join('ppt/slides', target)).replace('\\', '/')
            z.writestr('ppt/' + target.replace('../', ''), src.read(inner))
    os.remove(tmp)
    return out
